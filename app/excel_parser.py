import os
import json
import hashlib
import openpyxl
from typing import Dict, Any, List, Tuple, Optional
from datetime import datetime, date

# Verified Column Mappings from Real Excel Inspection
EXPECTED_HEADERS = {
    "Day Book": {
        "identifying_headers": ["Voucher Number", "Transaction Site", "Voucher Type"],
        "required_columns": [
            "Transaction Site", "Voucher Number", "Voucher Date", "Voucher Type",
            "Voucher Status", "Party Code", "Party Description", "Account Description",
            "Narration", "Created By", "Approved By"
        ]
    },
    "AP": {
        "identifying_headers": ["Accounting Site Code", "Voucher Number", "Voucher Type", "Invoice Number", "Total Voucher Amount"],
        "required_columns": [
            "Accounting Site Code", "Voucher Number", "Voucher Type", "Voucher Sub Type",
            "Party GST TIN", "Invoice Number", "Invoice Date", "Due Date",
            "Item/Service Description", "Item/Service Expense Account Description",
            "Booked Item Quantity", "Item/Service Rate", "Item/Service Detail Amount",
            "Total Tax Amount", "Total Voucher Amount", "Total TDS"
        ]
    },
    "AR": {
        "identifying_headers": ["Accounting Site Code", "Voucher Number", "Voucher Type", "Net Amount", "Total CGST", "Total SGST", "Total IGST"],
        "required_columns": [
            "Accounting Site Code", "Voucher Number", "Voucher Type", "Voucher Sub Type",
            "Item/Service Description", "Item Qty", "Item/Service Rate",
            "Item/Service Amount", "Item/Service Charges", "Item Amount Net Off Discount",
            "Item/Service Taxes", "Net Amount", "Total CGST", "Total SGST", "Total IGST"
        ]
    }
}

def format_cell_value(val: Any) -> Any:
    """Safely convert Excel cell values into clean Python types"""
    if val is None:
        return None
    if isinstance(val, (datetime, date)):
        return val.strftime("%d/%m/%Y")
    if isinstance(val, float):
        if val.is_integer():
            return int(val)
        return round(val, 4)
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ["nan", "none", "null"]:
        return None
    return val_str

def compute_hash(data: Dict[str, Any]) -> str:
    """Compute sha256 hash of a record's business keys and values for deduplication"""
    sorted_str = json.dumps(data, sort_keys=True, default=str)
    return hashlib.sha256(sorted_str.encode('utf-8')).hexdigest()

def detect_file_type_from_headers(col_name_to_idx: Dict[str, int]) -> str:
    """
    Deterministically identify whether an Excel file is Day Book, AP, or AR
    by evaluating its actual column signatures, completely independent of filename.
    """
    col_names_lower = [k.lower().strip() for k in col_name_to_idx.keys()]

    # Score Day Book signature
    daybook_score = 0
    if any("party code" in c for c in col_names_lower): daybook_score += 4
    if any("account description" in c for c in col_names_lower): daybook_score += 4
    if any(c == "transaction site" for c in col_names_lower): daybook_score += 3
    if any("approved by" in c for c in col_names_lower): daybook_score += 2
    if any("invoice number" in c for c in col_names_lower): daybook_score -= 10
    if any("booked item quantity" in c for c in col_names_lower): daybook_score -= 10

    # Score AP signature
    ap_score = 0
    if any("invoice number" in c or "invoice no" in c for c in col_names_lower): ap_score += 5
    if any("due date" in c for c in col_names_lower): ap_score += 4
    if any("total tds" in c for c in col_names_lower): ap_score += 4
    if any("party gst tin" in c or "party gst-tin" in c for c in col_names_lower): ap_score += 3
    if any("total voucher amount" in c for c in col_names_lower): ap_score += 3
    if any("net off discount" in c for c in col_names_lower): ap_score -= 10

    # Score AR signature
    ar_score = 0
    if any("net off discount" in c for c in col_names_lower): ar_score += 5
    if any("item/service taxes" in c or "item taxes" in c for c in col_names_lower): ar_score += 4
    if any("total cgst" in c or "total cost" in c for c in col_names_lower): ar_score += 3
    if any("total sgst" in c for c in col_names_lower): ar_score += 3
    if any("total igst" in c for c in col_names_lower): ar_score += 3
    if any("invoice number" in c for c in col_names_lower): ar_score -= 10

    scores = {"Day Book": daybook_score, "AP": ap_score, "AR": ar_score}
    best_type, best_score = max(scores.items(), key=lambda x: x[1])
    if best_score >= 4:
        return best_type
    return "Unknown"

def find_header_row_fast(rows_data: List[Tuple], max_search_rows: int = 30) -> Tuple[int, Dict[int, str], Dict[str, int]]:
    """
    Fast in-memory header detector from pre-loaded tuple rows.
    Returns: (header_row_idx_0_based, col_idx_to_name, col_name_to_idx)
    """
    best_row_idx = 0
    best_score = -1
    best_headers = {}

    search_limit = min(max_search_rows, len(rows_data))
    for r_idx in range(search_limit):
        row = rows_data[r_idx]
        if not row:
            continue
        row_headers = {}
        voucher_found = False
        non_empty = 0
        for c_idx, val in enumerate(row):
            if val is not None and str(val).strip():
                clean_name = str(val).strip()
                row_headers[c_idx] = clean_name
                non_empty += 1
                if "voucher number" in clean_name.lower() or "voucher no" in clean_name.lower() or "voucher type" in clean_name.lower():
                    voucher_found = True

        score = non_empty + (100 if voucher_found else 0)
        if score > best_score and non_empty >= 4:
            best_score = score
            best_row_idx = r_idx
            best_headers = row_headers

    name_to_idx = {name.strip(): idx for idx, name in best_headers.items()}
    return best_row_idx, best_headers, name_to_idx

def detect_file_type(filepath: str) -> str:
    """Detect whether an Excel file is Day Book, AP, or AR by inspecting internal header columns"""
    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        sheet = wb.active
        # Read first 30 rows only for ultra-fast detection
        sample_rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            sample_rows.append(row)
            if i >= 30:
                break
        wb.close()
        header_row_idx, col_idx_to_name, col_name_to_idx = find_header_row_fast(sample_rows)
        return detect_file_type_from_headers(col_name_to_idx)
    except Exception:
        return "Unknown"

def find_header_row(sheet, max_search_rows: int = 30) -> Tuple[int, Dict[int, str], Dict[str, int]]:
    """Legacy helper for backward compatibility"""
    rows_data = []
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        rows_data.append(row)
        if i >= max_search_rows:
            break
    r_idx, headers, name_to_idx = find_header_row_fast(rows_data, max_search_rows)
    # Convert 0-based idx to 1-based for legacy sheet.cell callers
    headers_1based = {c + 1: name for c, name in headers.items()}
    name_to_idx_1based = {name: c + 1 for name, c in name_to_idx.items()}
    return r_idx + 1, headers_1based, name_to_idx_1based

def validate_excel_file(filepath: str, expected_type: Optional[str] = None) -> Dict[str, Any]:
    """
    Validate Excel file structure and data presence against verified specifications,
    completely independent of filename with 2X fast read_only mode.
    """
    filename = os.path.basename(filepath)
    file_size = os.path.getsize(filepath)

    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        sheet = wb.active
        
        all_rows = list(sheet.iter_rows(values_only=True))
        wb.close()

        if not all_rows:
            return {
                "file_type": "Unknown",
                "filename": filename,
                "file_size_bytes": file_size,
                "detected_header_row": 0,
                "total_rows_detected": 0,
                "valid_data_rows": 0,
                "columns_matched": 0,
                "columns_missing": ["Empty spreadsheet"],
                "unexpected_columns": [],
                "status": "Failed",
                "message": f"File '{filename}' is empty."
            }

        header_row_idx, col_idx_to_name, col_name_to_idx = find_header_row_fast(all_rows)
        detected_type = expected_type or detect_file_type_from_headers(col_name_to_idx)
        
        if detected_type not in EXPECTED_HEADERS:
            return {
                "file_type": "Unknown",
                "filename": filename,
                "file_size_bytes": file_size,
                "detected_header_row": header_row_idx + 1,
                "total_rows_detected": len(all_rows),
                "valid_data_rows": 0,
                "columns_matched": 0,
                "columns_missing": ["Document columns do not match Day Book, AP, or AR specifications"],
                "unexpected_columns": [],
                "status": "Failed",
                "message": f"Could not detect Day Book, AP, or AR column structure in '{filename}'. Check column headers."
            }

        req_spec = EXPECTED_HEADERS[detected_type]
        missing_cols = []
        for req in req_spec["required_columns"]:
            matched = any(req.lower() == h.lower() or req.lower() in h.lower() for h in col_name_to_idx.keys())
            if not matched:
                missing_cols.append(req)

        # Fast row scan
        total_rows = 0
        valid_rows = 0
        v_idx = col_name_to_idx.get("Voucher Number") or col_name_to_idx.get("Voucher No.") or col_name_to_idx.get("Voucher No")

        for r_idx in range(header_row_idx + 1, len(all_rows)):
            row = all_rows[r_idx]
            if not row:
                continue
            v_val = row[v_idx] if (v_idx is not None and v_idx < len(row)) else row[0]
            
            if v_val is not None:
                v_str = str(v_val).strip()
                if v_str and not v_str.lower().startswith("exported on") and v_str.lower() != "total":
                    valid_rows += 1
            total_rows += 1

        has_missing_cols = len(missing_cols) > 0
        has_no_data = valid_rows == 0
        
        if has_missing_cols:
            status = "Failed"
            msg = f"Missing required columns: {', '.join(missing_cols)}"
        elif has_no_data:
            status = "Failed"
            msg = "File has valid headers but contains 0 valid data rows."
        else:
            status = "Valid"
            msg = f"Identified as {detected_type} with {valid_rows} verified data rows and all required columns."

        return {
            "file_type": detected_type,
            "filename": filename,
            "file_size_bytes": file_size,
            "detected_header_row": header_row_idx + 1,
            "total_rows_detected": total_rows,
            "valid_data_rows": valid_rows,
            "columns_matched": len(req_spec["required_columns"]) - len(missing_cols),
            "columns_missing": missing_cols,
            "unexpected_columns": [],
            "status": status,
            "message": msg
        }
    except Exception as e:
        return {
            "file_type": "Unknown",
            "filename": filename,
            "file_size_bytes": file_size,
            "detected_header_row": 0,
            "total_rows_detected": 0,
            "valid_data_rows": 0,
            "columns_matched": 0,
            "columns_missing": [str(e)],
            "unexpected_columns": [],
            "status": "Failed",
            "message": f"Error parsing Excel file: {str(e)}"
        }

def parse_daybook_file(filepath: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Parse Day Book Excel file with 2X high-speed read_only tuple indexing + raw preservation"""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    sheet = wb.active
    all_rows = list(sheet.iter_rows(values_only=True))
    wb.close()

    header_row_idx, col_idx_to_name, col_name_to_idx = find_header_row_fast(all_rows)

    def find_idx(candidates: List[str]) -> Optional[int]:
        for c in candidates:
            for actual_name, idx in col_name_to_idx.items():
                if c.lower() == actual_name.lower():
                    return idx
        return None

    idx_vno = find_idx(["Voucher Number", "Voucher No.", "Voucher No"])
    idx_site = find_idx(["Transaction Site", "Accounting Site Code", "Site"])
    idx_vdate = find_idx(["Voucher Date"])
    idx_vtype = find_idx(["Voucher Type"])
    idx_vstatus = find_idx(["Voucher Status"])
    idx_pcode = find_idx(["Party Code"])
    idx_pdesc = find_idx(["Party Description"])
    idx_adesc = find_idx(["Account Description"])
    idx_narr = find_idx(["Narration", "Header Narration", "Detail Narration"])
    idx_created_by = find_idx(["Created By"])
    idx_approved_by = find_idx(["Approved By"])
    idx_created_date = find_idx(["Created Date"])
    idx_approved_date = find_idx(["Approved Date"])

    def get_val(row: Tuple, idx: Optional[int]) -> Any:
        if idx is None or idx >= len(row):
            return None
        return format_cell_value(row[idx])

    records = []
    for r_idx in range(header_row_idx + 1, len(all_rows)):
        row = all_rows[r_idx]
        if not row:
            continue
        v_no = get_val(row, idx_vno)
        if not v_no or str(v_no).strip().lower() == "total" or str(v_no).lower().startswith("exported on"):
            continue

        raw_row = {}
        for c_idx, hname in col_idx_to_name.items():
            if c_idx < len(row):
                raw_row[hname] = format_cell_value(row[c_idx])

        site_val = get_val(row, idx_site)
        vdate_val = get_val(row, idx_vdate)
        vtype_val = get_val(row, idx_vtype)
        vstat_val = get_val(row, idx_vstatus)
        pcode_val = get_val(row, idx_pcode)
        pdesc_val = get_val(row, idx_pdesc)
        adesc_val = get_val(row, idx_adesc)
        narr_val = get_val(row, idx_narr)
        cby_val = get_val(row, idx_created_by)
        aby_val = get_val(row, idx_approved_by)
        cdate_val = get_val(row, idx_created_date)
        adate_val = get_val(row, idx_approved_date)

        record_dict = {
            "transaction_site": str(site_val) if site_val else None,
            "voucher_number": str(v_no),
            "voucher_date": str(vdate_val) if vdate_val else None,
            "voucher_type": str(vtype_val) if vtype_val else None,
            "voucher_status": str(vstat_val) if vstat_val else None,
            "party_code": str(pcode_val) if pcode_val else None,
            "party_description": str(pdesc_val) if pdesc_val else None,
            "account_description": str(adesc_val) if adesc_val else None,
            "narration": str(narr_val) if narr_val else None,
            "created_by": str(cby_val) if cby_val else None,
            "approved_by": str(aby_val) if aby_val else None,
            "created_date_raw": str(cdate_val) if cdate_val else None,
            "approved_date_raw": str(adate_val) if adate_val else None,
            "raw_data_json": json.dumps(raw_row, default=str),
            "row_hash": compute_hash(raw_row)
        }
        records.append(record_dict)

    stats = {"total_rows": len(records), "valid_rows": len(records), "header_row": header_row_idx + 1}
    return records, stats

def parse_ap_file(filepath: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Parse AP Excel file with 2X high-speed read_only tuple indexing + raw preservation"""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    sheet = wb.active
    all_rows = list(sheet.iter_rows(values_only=True))
    wb.close()

    header_row_idx, col_idx_to_name, col_name_to_idx = find_header_row_fast(all_rows)

    def find_idx(candidates: List[str]) -> Optional[int]:
        for c in candidates:
            for actual_name, idx in col_name_to_idx.items():
                if c.lower() == actual_name.lower():
                    return idx
        return None

    idx_vno = find_idx(["Voucher Number", "Voucher No.", "Voucher No"])
    idx_site = find_idx(["Accounting Site Code", "Transaction Site Code", "Transaction Site"])
    idx_vtype = find_idx(["Voucher Type"])
    idx_vsub = find_idx(["Voucher Sub Type", "Voucher Sub-Type"])
    idx_party_gst = find_idx(["Party GST TIN", "Party GST-TIN"])
    idx_inv_no = find_idx(["Invoice Number", "Invoice No.", "Invoice No"])
    idx_inv_date = find_idx(["Invoice Date"])
    idx_due_date = find_idx(["Due Date"])
    idx_item_desc = find_idx(["Item/Service Description", "Item/Services Description"])
    idx_exp_desc = find_idx(["Item/Service Expense Account Description", "Expense Account Description"])
    idx_qty = find_idx(["Booked Item Quantity", "Booked Quantity", "Item Qty"])
    idx_rate = find_idx(["Item/Service Rate", "Item Rate"])
    idx_detail_amt = find_idx(["Item/Service Detail Amount", "Item Detail Amount", "Item/Services Detail Amount"])
    idx_tax_amt = find_idx(["Total Tax Amount", "Tax Amount"])
    idx_v_amt = find_idx(["Total Voucher Amount", "Voucher Amount"])
    idx_tds = find_idx(["Total TDS", "TDS Amount"])
    idx_hnarr = find_idx(["Header Narration", "Narration"])
    idx_dnarr = find_idx(["Detail Narration", "Narration/Remarks"])

    def get_val(row: Tuple, idx: Optional[int]) -> Any:
        if idx is None or idx >= len(row):
            return None
        return format_cell_value(row[idx])

    def get_float(row: Tuple, idx: Optional[int]) -> Optional[float]:
        val = get_val(row, idx)
        if val is None:
            return None
        try:
            return float(val)
        except (ValueError, TypeError):
            return None

    records = []
    for r_idx in range(header_row_idx + 1, len(all_rows)):
        row = all_rows[r_idx]
        if not row:
            continue
        v_no = get_val(row, idx_vno)
        if not v_no or str(v_no).strip().lower() == "total" or str(v_no).lower().startswith("exported on"):
            continue

        raw_row = {}
        for c_idx, hname in col_idx_to_name.items():
            if c_idx < len(row):
                raw_row[hname] = format_cell_value(row[c_idx])

        site_val = get_val(row, idx_site)
        vtype_val = get_val(row, idx_vtype)
        vsub_val = get_val(row, idx_vsub)
        pgst_val = get_val(row, idx_party_gst)
        invno_val = get_val(row, idx_inv_no)
        invdate_val = get_val(row, idx_inv_date)
        duedate_val = get_val(row, idx_due_date)
        itemdesc_val = get_val(row, idx_item_desc)
        expdesc_val = get_val(row, idx_exp_desc)
        hnarr_val = get_val(row, idx_hnarr)
        dnarr_val = get_val(row, idx_dnarr)

        record_dict = {
            "accounting_site_code": str(site_val) if site_val else None,
            "voucher_number": str(v_no),
            "voucher_type": str(vtype_val) if vtype_val else None,
            "voucher_sub_type": str(vsub_val) if vsub_val else None,
            "party_gst_tin": str(pgst_val) if pgst_val else None,
            "invoice_number": str(invno_val) if invno_val else None,
            "invoice_date": str(invdate_val) if invdate_val else None,
            "due_date": str(duedate_val) if duedate_val else None,
            "item_service_description": str(itemdesc_val) if itemdesc_val else None,
            "item_service_expense_account_desc": str(expdesc_val) if expdesc_val else None,
            "booked_item_quantity": get_float(row, idx_qty),
            "item_service_rate": get_float(row, idx_rate),
            "item_service_detail_amount": get_float(row, idx_detail_amt),
            "total_tax_amount": get_float(row, idx_tax_amt),
            "total_voucher_amount": get_float(row, idx_v_amt),
            "total_tds": get_float(row, idx_tds),
            "header_narration": str(hnarr_val) if hnarr_val else None,
            "detail_narration": str(dnarr_val) if dnarr_val else None,
            "raw_data_json": json.dumps(raw_row, default=str),
            "row_hash": compute_hash(raw_row)
        }
        records.append(record_dict)

    stats = {"total_rows": len(records), "valid_rows": len(records), "header_row": header_row_idx + 1}
    return records, stats

def parse_ar_file(filepath: str) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Parse AR Excel file with 2X high-speed read_only tuple indexing + raw preservation"""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    sheet = wb.active
    all_rows = list(sheet.iter_rows(values_only=True))
    wb.close()

    header_row_idx, col_idx_to_name, col_name_to_idx = find_header_row_fast(all_rows)

    def find_idx(candidates: List[str]) -> Optional[int]:
        for c in candidates:
            for actual_name, idx in col_name_to_idx.items():
                if c.lower() == actual_name.lower():
                    return idx
        return None

    idx_vno = find_idx(["Voucher Number", "Voucher No.", "Voucher No"])
    idx_site = find_idx(["Accounting Site Code", "Transaction Site Code", "Transaction Site"])
    idx_vtype = find_idx(["Voucher Type"])
    idx_vsub = find_idx(["Voucher Sub Type", "Voucher Sub-Type"])
    idx_item_desc = find_idx(["Item/Service Description", "Item/Services Description"])
    idx_qty = find_idx(["Item Qty", "Item Quantity", "Booked Item Quantity"])
    idx_rate = find_idx(["Item/Service Rate", "Item Rate"])
    idx_amt = find_idx(["Item/Service Amount", "Item Amount"])
    idx_charges = find_idx(["Item/Service Charges", "Item/Services Charge"])
    idx_net_off = find_idx(["Item Amount Net Off Discount", "Item/Amount Net Off Discount"])
    idx_taxes = find_idx(["Item/Service Taxes", "Item/Services Taxes"])
    idx_net_amt = find_idx(["Net Amount", "Total Voucher Amount"])
    idx_cgst = find_idx(["Total CGST", "Total COST"])
    idx_sgst = find_idx(["Total SGST"])
    idx_igst = find_idx(["Total IGST"])
    idx_remarks = find_idx(["Narration/Remarks", "Narration", "Detail Narration"])

    def get_val(row: Tuple, idx: Optional[int]) -> Any:
        if idx is None or idx >= len(row):
            return None
        return format_cell_value(row[idx])

    def get_float(row: Tuple, idx: Optional[int]) -> Optional[float]:
        val = get_val(row, idx)
        if val is None:
            return None
        try:
            return float(val)
        except (ValueError, TypeError):
            return None

    records = []
    for r_idx in range(header_row_idx + 1, len(all_rows)):
        row = all_rows[r_idx]
        if not row:
            continue
        v_no = get_val(row, idx_vno)
        if not v_no or str(v_no).strip().lower() == "total" or str(v_no).lower().startswith("exported on"):
            continue

        raw_row = {}
        for c_idx, hname in col_idx_to_name.items():
            if c_idx < len(row):
                raw_row[hname] = format_cell_value(row[c_idx])

        site_val = get_val(row, idx_site)
        vtype_val = get_val(row, idx_vtype)
        vsub_val = get_val(row, idx_vsub)
        itemdesc_val = get_val(row, idx_item_desc)
        remarks_val = get_val(row, idx_remarks)

        record_dict = {
            "accounting_site_code": str(site_val) if site_val else None,
            "voucher_number": str(v_no),
            "voucher_type": str(vtype_val) if vtype_val else None,
            "voucher_sub_type": str(vsub_val) if vsub_val else None,
            "item_service_description": str(itemdesc_val) if itemdesc_val else None,
            "item_quantity": get_float(row, idx_qty),
            "item_service_rate": get_float(row, idx_rate),
            "item_service_amount": get_float(row, idx_amt),
            "item_service_charges": get_float(row, idx_charges),
            "item_amount_net_off_discount": get_float(row, idx_net_off),
            "item_service_taxes": get_float(row, idx_taxes),
            "net_amount": get_float(row, idx_net_amt),
            "total_cgst": get_float(row, idx_cgst),
            "total_sgst": get_float(row, idx_sgst),
            "total_igst": get_float(row, idx_igst),
            "narration_remarks": str(remarks_val) if remarks_val else None,
            "raw_data_json": json.dumps(raw_row, default=str),
            "row_hash": compute_hash(raw_row)
        }
        records.append(record_dict)

    stats = {"total_rows": len(records), "valid_rows": len(records), "header_row": header_row_idx + 1}
    return records, stats


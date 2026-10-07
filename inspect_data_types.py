import os
import openpyxl

data_dir = r"d:\Shivam AI\Shivam AI\Day Book\data_files"

def inspect_types(filename, header_row, check_cols):
    wb = openpyxl.load_workbook(os.path.join(data_dir, filename), data_only=True)
    sheet = wb.active
    print(f"\n--- {filename} Types Check ---")
    headers = {c: sheet.cell(header_row, c).value for c in range(1, sheet.max_column + 1)}
    
    for col_idx in check_cols:
        hname = headers.get(col_idx)
        vals = [sheet.cell(r, col_idx).value for r in range(header_row + 1, min(header_row + 10, sheet.max_row + 1))]
        types = [type(v).__name__ for v in vals]
        print(f"Col {col_idx:3d} ({hname}): Sample Types: {types[:4]} | Values: {vals[:4]}")

# Day Book
inspect_types("DayBookReport_06-10-2026-14-27-32.xlsx", 13, [1, 6, 8, 11, 12, 13, 14, 19, 40, 41, 43])
# AP
inspect_types("APTaxRegisterDetailReport_06-10-2026-14-30-40.xlsx", 14, [1, 8, 14, 16, 44, 52, 53, 55, 68, 70, 75, 79, 80, 83, 84, 95])
# AR
inspect_types("ARTaxRegisterDetailReport_06-10-2026-14-58-25.xlsx", 15, [1, 6, 7, 11, 65, 72, 77, 84, 85, 87, 88, 90, 121, 137, 153])

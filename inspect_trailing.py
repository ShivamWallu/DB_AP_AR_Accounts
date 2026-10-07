import os
import openpyxl

data_dir = r"d:\Shivam AI\Shivam AI\Day Book\data_files"

for fname, hrow in [
    ("DayBookReport_06-10-2026-14-27-32.xlsx", 13),
    ("APTaxRegisterDetailReport_06-10-2026-14-30-40.xlsx", 14),
    ("ARTaxRegisterDetailReport_06-10-2026-14-58-25.xlsx", 15)
]:
    wb = openpyxl.load_workbook(os.path.join(data_dir, fname), data_only=True)
    sheet = wb.active
    print(f"\nLast 10 rows of {fname}:")
    for r in range(sheet.max_row - 9, sheet.max_row + 1):
        vals = [sheet.cell(r, c).value for c in range(1, min(15, sheet.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"  Row {r}: {vals[:6]}")

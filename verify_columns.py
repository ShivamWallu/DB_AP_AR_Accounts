import os
import openpyxl
from openpyxl.utils import get_column_letter

data_dir = r"d:\Shivam AI\Shivam AI\Day Book\data_files"

files_info = {
    "Day Book": {
        "file": "DayBookReport_06-10-2026-14-27-32.xlsx",
        "requested_common": {
            "F": "Voucher Number",
            "A": "Transaction Site",
            "K": "Voucher Type"
        },
        "requested_main": {
            "A": "Transaction Site",
            "F": "Voucher No.",
            "H": "Voucher Date",
            "K": "Voucher Type",
            "L": "Voucher Status",
            "M": "Party Code",
            "N": "Party Description",
            "S": "Account Description",
            "AN": "Narration",
            "AO": "Created By",
            "AQ": "Approved By"
        }
    },
    "AP": {
        "file": "APTaxRegisterDetailReport_06-10-2026-14-30-40.xlsx",
        "requested_common": {
            "H": "Voucher No.",
            "A": "Accounting Site Code",
            "N": "Voucher Type"
        },
        "requested_main": {
            "P": "Voucher Sub-Type",
            "AR": "Party GST TIN",
            "AZ": "Invoice No.",
            "BA": "Invoice Date",
            "BC": "Due Date",
            "BP": "Item/Services Description",
            "BR": "Item/Service Expense Account Description",
            "BW": "Booked Item Quantity",
            "CA": "Item/Service Rate",
            "CB": "Item/Services Detail Amount",
            "CE": "Total Tax Amount",
            "CF": "Total Voucher Amount",
            "CQ": "Total TDS"
        }
    },
    "AR": {
        "file": "ARTaxRegisterDetailReport_06-10-2026-14-58-25.xlsx",
        "requested_common": {
            "A": "Account Site Code",
            "K": "Voucher No.",
            "F": "Voucher Type"
        },
        "requested_main": {
            "G": "Voucher Sub-Type",
            "BM": "Item/Services Description",
            "BT": "Item Quantity",
            "BY": "Item/Services Rate",
            "CF": "Item/Services Amount",
            "CG": "Item/Services Charge",
            "CI": "Item/Amount Net Off Discount",
            "CJ": "Item/Services Taxes",
            "CL": "Net Amount",
            "DQ": "Total COST", # Let's see if this is Total CGST or Total COST
            "EH": "Total SGST", # Let's see
            "EW": "Total IGST"  # Let's see
        }
    }
}

for doc_type, info in files_info.items():
    filepath = os.path.join(data_dir, info["file"])
    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet = wb.active
    
    # Find header row
    best_row = 1
    max_cols = 0
    for r in range(1, min(35, sheet.max_row + 1)):
        non_empty = sum(1 for c in range(1, sheet.max_column + 1) if sheet.cell(r, c).value is not None)
        if non_empty > max_cols:
            max_cols = non_empty
            best_row = r
            
    print(f"\n{'='*90}")
    print(f"DOCUMENT TYPE: {doc_type}")
    print(f"File: {info['file']}")
    print(f"Header Row Identified: Row {best_row} (Total non-empty headers: {max_cols})")
    print(f"{'='*90}")
    
    # Extract actual headers
    actual_headers = {}
    for c in range(1, sheet.max_column + 1):
        col_letter = get_column_letter(c)
        val = sheet.cell(best_row, c).value
        actual_headers[col_letter] = (c, str(val).strip() if val is not None else "")
        
    print("\n--- VERIFYING REQUESTED COMMON COLUMNS ---")
    for letter, expected_name in info["requested_common"].items():
        actual = actual_headers.get(letter, (None, "[COLUMN NOT FOUND]"))
        match_status = "MATCH" if expected_name.lower() in actual[1].lower() or actual[1].lower() in expected_name.lower() else "MISMATCH / DIFF"
        print(f"Col {letter:4s} | Expected: '{expected_name}' | Actual Header: '{actual[1]}' | Status: {match_status}")
        
    print("\n--- VERIFYING REQUESTED MAIN COLUMNS ---")
    for letter, expected_name in info["requested_main"].items():
        actual = actual_headers.get(letter, (None, "[COLUMN NOT FOUND]"))
        match_status = "MATCH" if expected_name.lower() in actual[1].lower() or actual[1].lower() in expected_name.lower() else "MISMATCH / DIFF"
        print(f"Col {letter:4s} | Expected: '{expected_name}' | Actual Header: '{actual[1]}' | Status: {match_status}")

    # Check for Narration, Created By, Approved By across the file
    print("\n--- SEARCHING FOR Narration, Created By, Approved By in this file ---")
    found_special = {}
    for col_letter, (col_idx, hname) in actual_headers.items():
        for target in ["narration", "created", "approved", "prepared", "verified", "author"]:
            if target in hname.lower():
                found_special[col_letter] = (col_idx, hname)
    if found_special:
        for ltr, (idx, hnm) in found_special.items():
            print(f"  Found potential column {ltr} (col {idx}): '{hnm}'")
    else:
        print("  None found in this file.")

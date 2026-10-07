import os
import openpyxl
from openpyxl.utils import get_column_letter

data_dir = r"d:\Shivam AI\Shivam AI\Day Book\data_files"
wb = openpyxl.load_workbook(os.path.join(data_dir, "ARTaxRegisterDetailReport_06-10-2026-14-58-25.xlsx"), data_only=True)
sheet = wb.active

for col_idx in range(110, 160):
    letter = get_column_letter(col_idx)
    header = sheet.cell(15, col_idx).value
    val = sheet.cell(16, col_idx).value
    if header and ("total" in header.lower() or "gst" in header.lower() or "cost" in header.lower()):
        print(f"Col {letter:4s} ({col_idx:3d}): '{header}' | Sample: {val}")
